#!/usr/bin/env python3
"""Bangun sheets/gmgn-sheet-template.xlsx dari sheets/csv/*.csv.

CSV adalah sumber kebenaran. Kolom ID dibuat berformat Plain Text agar ID Discord
(18-19 digit) tidak dibulatkan oleh Google Sheets / Excel.

Pemakaian:  python scripts/make_template.py
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
CSV_DIR = ROOT / "sheets" / "csv"
OUT = ROOT / "sheets" / "gmgn-sheet-template.xlsx"

FONT = "Arial"
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
BODY = Font(name=FONT, size=10)
EDIT_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# nama sheet -> (lebar kolom, indeks kolom (1-based) yang wajib Plain Text)
LAYOUT = {
    "Config": ([24, 32, 80], (2,)),
    "Channels": ([26, 10, 10, 10], (1, 3, 4)),
    "Messages": ([10, 40, 10], ()),
    "Logs": ([20, 28, 24, 26, 8, 30, 26, 12, 50], (4, 7)),
}

PETUNJUK = [
    ("TEMPLATE SHEET - GM/GN AI AGENT (n8n)", True),
    ("", False),
    ("Cara pakai: Google Sheets > File > Import > Upload file ini > 'Create new spreadsheet'.", False),
    ("Lalu: File > Save as Google Sheets bila di judul masih ada label .XLSX (n8n tidak bisa membaca file Office).", False),
    ("Jangan ganti nama sheet maupun nama kolom (workflow n8n memakai nama persis ini).", False),
    ("", False),
    ("Config   : isi sel kuning di kolom B (monitor_channel_id, admin_user_id, monitor_webhook_url). gn_time = jam GN universal. last_message_id biarkan kosong.", False),
    ("            monitor_webhook_url (opsional): URL webhook channel pemantau agar bukti dan balasan tampil sebagai kartu embed.", False),
    ("Channels : channel_id, aktif (TRUE/FALSE), gm_time, gn_time (format HH:mm). Scheduler harian saat ini memakai Config.gn_time; kolom jam dicatat untuk referensi.", False),
    ("            Nama server dan nama channel diambil otomatis dari Discord, tidak perlu ditulis.", False),
    ("Messages : daftar pesan per tipe (gm/gn). bobot = peluang relatif, makin besar makin sering muncul.", False),
    ("Logs     : diisi otomatis oleh n8n (bukti kirim). Jangan diedit manual.", False),
    ("", False),
    ("PENTING: kolom ID dan jam sudah berformat Plain Text. Jangan ubah ke Number: ID Discord 18-19 digit", False),
    ("akan dibulatkan dan jadi salah. Setelah import, cek Format > Number > Plain text.", False),
    ("", False),
    ("Legenda: sel kuning = wajib kamu isi sendiri.", False),
]


def read_csv(name):
    with open(CSV_DIR / f"{name}.csv", newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Petunjuk"
    for i, (text, bold) in enumerate(PETUNJUK, 1):
        ws.cell(row=i, column=1, value=text).font = Font(name=FONT, size=12 if bold else 10, bold=bold)
    ws.column_dimensions["A"].width = 120

    for name, (widths, text_cols) in LAYOUT.items():
        rows = read_csv(name)
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=value if value != "" else None)
                if r == 1:
                    cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font, cell.border = BODY, BORDER
        for c in text_cols:
            for r in range(2, 501):
                ws.cell(row=r, column=c).number_format = "@"
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22

    wb["Config"]["B2"].fill = EDIT_FILL
    wb["Config"]["B3"].fill = EDIT_FILL
    wb["Config"]["B4"].fill = EDIT_FILL

    dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    wb["Channels"].add_data_validation(dv)
    dv.add("B2:B500")
    dv2 = DataValidation(type="list", formula1='"gm,gn"', allow_blank=True)
    wb["Messages"].add_data_validation(dv2)
    dv2.add("A2:A500")

    wb.save(OUT)
    print(f"OK -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    sys.exit(main())
